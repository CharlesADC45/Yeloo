from pathlib import Path
from datetime import date
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\oluac\Desktop\YELOO+")
SOURCE = Path(r"C:\Users\oluac\Downloads\WEB")
OUT = ROOT / "documentation" / "livrables_adc"
TMP = ROOT / ".tmp" / "adc_documentation"
OUT.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)

DOCX = OUT / "Documentation_fonctionnelle_et_UI_ADC.docx"

NAVY = "0E2A3F"
BLUE = "24558E"
CYAN = "42BEE3"
PALE = "EAF6FB"
LIGHT = "F4F7FA"
MID = "627487"
WHITE = "FFFFFF"
INK = "172B3A"
YELLOW = "FFC928"

IMAGES = {
    "home": SOURCE / "01 — Accueil · Version principale.png",
    "home_alt": SOURCE / "01B — Accueil · Proposition alternative.png",
    "home_classic": SOURCE / "01C — Accueil · Proposition classique.png",
    "contact": SOURCE / "02 — Contact.png",
    "solutions": SOURCE / "03 — Solutions · Vue principale.png",
    "solutions_hover": SOURCE / "03A — Solutions · État hover.png",
    "about": SOURCE / "04 — À propos d’ADC.png",
    "former": SOURCE / "05A — Notre approche · Former.png",
    "equiper": SOURCE / "05B — Notre approche · Équiper.png",
    "accompagner": SOURCE / "05C — Notre approche · Accompagner.png",
    "careers_list": SOURCE / "06 — Carrières · Liste des postes.png",
    "careers_detail": SOURCE / "06 — Carrières · Offres et détail.png",
    "cases": SOURCE / "07 — Études de cas · Détails.png",
    "concept": SOURCE / "CONCEPT — Hero et première impression.png",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def font(run, size=10.5, color=INK, bold=False, italic=False, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def add_text(doc, text, size=10.5, color=INK, bold=False, italic=False,
             align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_with_next = keep
    font(p.add_run(text), size, color, bold, italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    font(p.add_run(text), 10.1, INK)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def kicker(doc, text):
    p = add_text(doc, text.upper(), 8.5, CYAN, True, after=3, keep=True)
    p.paragraph_format.letter_spacing = Pt(0.6) if hasattr(p.paragraph_format, "letter_spacing") else None
    return p


def callout(doc, title, body, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.25)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 140, 180, 140, 180)
    remove_table_borders(table)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(title), 10.5, BLUE, True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.08
    font(p2.add_run(body), 9.7, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def image_copy(path, max_px=1800):
    out = TMP / (path.stem.replace(" ", "_") + "_doc.jpg")
    with Image.open(path) as im:
        im = im.convert("RGB")
        scale = min(1.0, max_px / max(im.size))
        if scale < 1:
            im = im.resize((int(im.width * scale), int(im.height * scale)), Image.Resampling.LANCZOS)
        im.save(out, "JPEG", quality=88, optimize=True)
    return out


def add_image(doc, key, width=None, height=None, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    img = image_copy(IMAGES[key])
    kw = {}
    if width: kw["width"] = Inches(width)
    if height: kw["height"] = Inches(height)
    p.add_run().add_picture(str(img), **kw)
    if caption:
        cp = add_text(doc, caption, 8.2, MID, italic=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER, after=7)
        cp.paragraph_format.keep_with_next = False


def add_two_images(doc, items, width=2.85, height=None):
    table = doc.add_table(rows=1, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    for i, (key, label) in enumerate(items):
        cell = table.cell(0, i)
        set_cell_margins(cell, 60, 90, 60, 90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_args = {"height": Inches(height)} if height else {"width": Inches(width)}
        p.add_run().add_picture(str(image_copy(IMAGES[key])), **picture_args)
        c = cell.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(0)
        font(c.add_run(label), 8.5, BLUE, True)


def page_title(doc, section_no, title, subtitle=None):
    kicker(doc, f"{section_no} / Documentation ADC")
    heading(doc, title, 1)
    if subtitle:
        add_text(doc, subtitle, 10.5, MID, after=10)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(3)
    font(p.add_run("AFRICA DIGITAL CONNECT  |  Documentation du site web  |  2026"), 8, MID)


def set_doc_styles(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.30)
    sec.footer_distance = Inches(0.30)
    add_footer(sec)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in [
        ("Heading 1", 21, NAVY, 4, 9),
        ("Heading 2", 14, BLUE, 10, 5),
        ("Heading 3", 11.5, NAVY, 7, 3),
    ]:
        st = doc.styles[name]
        st.font.name = "Aptos Display"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2"):
        st = doc.styles[name]
        st.font.name = "Aptos"
        st.font.size = Pt(10.1)
        st.paragraph_format.space_after = Pt(4)


def add_page_break(doc):
    doc.add_page_break()


def add_matrix(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for i, w in enumerate(widths):
            table.columns[i].width = Inches(w)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        font(p.add_run(value), 8.8, WHITE, True)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            if ridx % 2 == 1:
                set_cell_shading(cells[i], LIGHT)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            font(p.add_run(str(value)), 8.7, INK, bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def build():
    doc = Document()
    set_doc_styles(doc)

    # Cover
    add_text(doc, "AFRICA DIGITAL CONNECT", 10, CYAN, True, after=45)
    add_text(doc, "Documentation\nfonctionnelle et UI", 32, NAVY, True, after=10)
    add_text(doc, "Site web institutionnel ADC", 17, BLUE, True, after=16)
    add_text(doc, "Référentiel de conception, de contenu et de recette", 11.5, MID, after=26)
    add_image(doc, "home", height=4.25)
    add_text(doc, "Version 1.0  |  21 juillet 2026", 9, MID, align=WD_ALIGN_PARAGRAPH.RIGHT, before=10)

    add_page_break(doc)
    page_title(doc, "01", "Objet et périmètre", "Une base commune pour cadrer la réalisation et valider l’expérience.")
    callout(doc, "Finalité du document", "Formaliser les intentions visibles dans les maquettes, harmoniser les pages, préciser les comportements attendus et fournir une checklist de recette avant mise en production.")
    heading(doc, "Périmètre couvert", 2)
    for item in [
        "Accueil : trois propositions et un concept de première impression.",
        "Pages institutionnelles : À propos, Notre approche, Contact.",
        "Offre : Solutions et état de survol des cartes.",
        "Preuves : études de cas détaillées.",
        "Recrutement : liste des postes et fiche d’offre.",
        "Composants transverses : navigation, appels à l’action, formulaires, partenaires et pied de page.",
    ]: add_bullet(doc, item)
    heading(doc, "Principes directeurs", 2)
    add_matrix(doc, ["Principe", "Traduction dans l’interface"], [
        ("Clarté", "Une promesse simple, un titre fort et une action principale par zone."),
        ("Crédibilité", "Chiffres clés, références partenaires, cas clients et expertise visibles."),
        ("Proximité", "Photographies humaines, ton direct et accès rapide au contact."),
        ("Cohérence", "Même navigation, même palette et mêmes composants sur l’ensemble du site."),
        ("Conversion", "CTA “Demander une démo” persistant et formulaires courts."),
    ], [1.35, 4.9])

    add_page_break(doc)
    page_title(doc, "02", "Architecture de l’information")
    add_matrix(doc, ["Niveau 1", "Pages / contenus", "Objectif utilisateur"], [
        ("Accueil", "Promesse, offres PME, solutions, approche, cas, carrières, contact", "Comprendre ADC et choisir la prochaine action"),
        ("Qui sommes-nous ?", "Présentation, valeurs, ambitions, RSE, équipe", "Évaluer la légitimité et la culture"),
        ("Notre approche", "Former, Équiper, Accompagner", "Identifier le mode d’intervention adapté"),
        ("Nos solutions", "Balafon, CCR, Baobab, Bamboo", "Explorer les produits et demander une démo"),
        ("Études de cas", "Défi, réponse ADC, impact", "Vérifier les résultats en contexte réel"),
        ("Carrières", "Liste, détail, candidature", "Trouver une offre et postuler"),
        ("Contact", "Formulaire, coordonnées, carte", "Échanger avec l’équipe"),
    ], [1.2, 2.9, 2.15])
    heading(doc, "Navigation recommandée", 2)
    for item in [
        "En-tête fixe ou réapparaissant au défilement, avec état actif clairement identifiable.",
        "CTA primaire “Demander une démo” contrasté et identique sur toutes les pages.",
        "Libellés homogènes : conserver “Nos solutions”, “Études de cas” et “Carrières” partout.",
        "Sur mobile, menu compact avec accès direct au contact et au CTA principal.",
        "Pied de page simplifié : supprimer les doublons et garantir des liens uniques et lisibles.",
    ]: add_bullet(doc, item)
    callout(doc, "Point de vigilance", "Plusieurs maquettes montrent des doublons ou chevauchements dans le pied de page. La version développée doit utiliser un seul composant partagé, alimenté par une source de contenu unique.", "FFF5D8")

    add_page_break(doc)
    page_title(doc, "03", "Système visuel et composants")
    heading(doc, "Palette observée", 2)
    add_matrix(doc, ["Rôle", "Couleur", "Usage recommandé"], [
        ("Bleu nuit", "#0E2A3F", "Hero, cartes fortes, surfaces de contraste"),
        ("Bleu ADC", "#24558E", "Titres, bandeaux, pied de page"),
        ("Cyan", "#42BEE3", "CTA, accents, libellés et liens"),
        ("Blanc", "#FFFFFF", "Fond principal et texte inversé"),
        ("Gris clair", "#F4F7FA", "Sections secondaires et cartes"),
        ("Jaune", "#FFC928", "Accent exceptionnel issu du symbole Afrique"),
    ], [1.25, 1.2, 3.8])
    heading(doc, "Composants réutilisables", 2)
    for item in [
        "Header : logo, navigation, contact avec pictogramme et CTA démo.",
        "Hero : sur-titre, titre H1, texte d’appui, visuel ou fond photographique.",
        "Carte : image, voile sombre, titre, description courte et état hover/focus.",
        "Bandeau chiffres : quatre indicateurs avec libellés contextualisés.",
        "Formulaire : champs étiquetés, validation, consentement et confirmation d’envoi.",
        "Footer : coordonnées, services, liens rapides, newsletter, réseaux et mentions légales.",
    ]: add_bullet(doc, item)
    heading(doc, "Règles de cohérence", 2)
    add_text(doc, "Limiter chaque écran à un CTA primaire. Les cartes interactives doivent avoir les mêmes rayons, hauteurs, ombres, comportements de clavier et transitions. Les titres doivent suivre une hiérarchie H1 > H2 > H3, sans utiliser la taille seule pour transmettre le niveau.")

    add_page_break(doc)
    page_title(doc, "04", "Accueil - version principale", "Version la plus complète et la plus adaptée à un lancement.")
    add_two_images(doc, [("home", "Version principale"), ("concept", "Concept hero / première impression")], height=3.75)
    heading(doc, "Lecture fonctionnelle", 2)
    for item in [
        "Le hero positionne ADC comme DSI des PME et donne immédiatement accès à la démo.",
        "Les chiffres clés renforcent la confiance, mais leurs libellés doivent être spécifiques et vérifiables.",
        "Le bloc “Nos offres pour les PME” rend l’offre concrète grâce à sept familles de services.",
        "Les solutions, l’approche, les études de cas et les carrières construisent une progression de preuve.",
        "Le contact en fin de page capte les visiteurs qui ont parcouru l’argumentaire complet.",
    ]: add_bullet(doc, item)
    callout(doc, "Recommandation", "Retenir la version principale comme base. Utiliser le hero de la proposition classique si l’on privilégie une première impression plus légère. Le pop-up du concept doit rester optionnel, temporisé, facilement fermable et non intrusif sur mobile.")

    add_page_break(doc)
    page_title(doc, "05", "Accueil - comparaison des variantes")
    add_two_images(doc, [("home_alt", "Proposition alternative"), ("home_classic", "Proposition classique")], height=3.65)
    add_matrix(doc, ["Critère", "Principale", "Alternative", "Classique"], [
        ("Lisibilité du hero", "Bonne", "Moyenne - fond chargé", "Très bonne"),
        ("Richesse éditoriale", "Très élevée", "Élevée", "Élevée"),
        ("Preuve par les cas", "Intégrée", "Non visible", "Non visible"),
        ("Carrières", "Intégrées", "Bandeau court", "Bandeau court"),
        ("Recommandation", "Base de production", "Variante campagne", "Source pour le hero"),
    ], [1.35, 1.55, 1.55, 1.55])
    heading(doc, "Décisions à figer", 2)
    for item in [
        "Choisir une seule version canonique pour éviter la divergence du contenu.",
        "Conserver un seul bouton principal dans le hero ; le second bouton observé dans une variante est redondant.",
        "Normaliser la casse des titres et la terminologie “PME”.",
        "Remplacer les chiffres génériques par des métriques réelles et datées.",
    ]: add_bullet(doc, item)

    add_page_break(doc)
    page_title(doc, "06", "Qui sommes-nous ?")
    add_image(doc, "about", height=5.35, caption="Maquette 04 - présentation institutionnelle d’Africa Digital Connect")
    heading(doc, "Contenu attendu", 2)
    add_text(doc, "La page regroupe l’histoire et la mission d’ADC, ses valeurs, ses ambitions, sa politique RSE et l’équipe dirigeante. Elle doit répondre à trois questions : qui est ADC, pourquoi lui faire confiance et qui porte l’engagement.")
    for item in [
        "Réécrire le texte en paragraphes courts et en casse normale pour améliorer la lecture.",
        "Associer chaque membre de l’équipe à un nom, une fonction et une biographie courte vérifiés.",
        "Documenter les engagements RSE avec des actions, dates et résultats mesurables.",
        "Ajouter des textes alternatifs pertinents aux photographies.",
    ]: add_bullet(doc, item)

    add_page_break(doc)
    page_title(doc, "07", "Notre approche : Former, Équiper, Accompagner")
    add_two_images(doc, [("former", "Former"), ("equiper", "Équiper"), ("accompagner", "Accompagner")], height=3.45)
    add_matrix(doc, ["Pilier", "Promesse", "Preuves / contenus"], [
        ("Former", "Développer les compétences numériques et métier", "Catalogue, objectifs, publics, modalités, témoignages"),
        ("Équiper", "Fournir des infrastructures et outils performants", "Expertises, projets, partenaires, technologies"),
        ("Accompagner", "Mobiliser des experts et accélérer les déploiements", "Services, délais, gouvernance, résultats"),
    ], [1.05, 2.25, 2.95])
    callout(doc, "À harmoniser", "Les cartes utilisent actuellement des libellés génériques répétés. Chaque carte doit porter un titre réel, une description courte et un lien vers un détail ou une prise de contact.")

    add_page_break(doc)
    page_title(doc, "08", "Nos solutions")
    add_two_images(doc, [("solutions", "Vue principale"), ("solutions_hover", "État de survol")], 2.85)
    heading(doc, "Comportement attendu des cartes", 2)
    for item in [
        "État repos : titre, visuel, contraste suffisant et indication d’interactivité.",
        "Hover : révéler la description ou l’action sans déplacer la mise en page.",
        "Focus clavier : reproduire l’information du hover avec un contour visible.",
        "Activation : ouvrir une page solution ou une fiche détaillée, jamais une zone sans destination.",
        "Mobile : afficher l’information essentielle sans dépendre du survol.",
    ]: add_bullet(doc, item)
    heading(doc, "Catalogue", 2)
    add_text(doc, "Balafon - supervision et monitoring IT ; Collaborateur Cyber Responsable (CCR) - centre de contrôle réseau et cybersécurité ; Baobab - gestion des services IT ; Bamboo - collaboration et productivité.")

    add_page_break(doc)
    page_title(doc, "09", "Études de cas")
    add_image(doc, "cases", height=5.45, caption="Maquette 07 - trois cas structurés par défi, réponse et impact")
    heading(doc, "Gabarit éditorial", 2)
    add_matrix(doc, ["Champ", "Exigence"], [
        ("Contexte", "Secteur, organisation et situation initiale sans divulguer d’information sensible."),
        ("Défi", "Problème précis, conséquences et contraintes."),
        ("Réponse ADC", "Solution, méthode, technologies et accompagnement."),
        ("Impact", "Résultats quantifiés, période de mesure et source de validation."),
        ("CTA", "Découvrir la solution associée ou demander un échange."),
    ], [1.35, 4.9])
    callout(doc, "Crédibilité", "Éviter les résultats vagues. Chaque impact doit être mesurable ou explicitement qualifié comme retour qualitatif.")

    add_page_break(doc)
    page_title(doc, "10", "Carrières - liste des postes")
    add_image(doc, "careers_list", height=5.35, caption="Maquette 06 - grille de six opportunités")
    heading(doc, "Fonctions clés", 2)
    for item in [
        "Afficher le nombre d’offres ouvertes et actualiser automatiquement l’état des postes.",
        "Permettre le filtrage par métier, type de contrat et localisation si le volume augmente.",
        "Chaque carte contient : contrat, intitulé, lieu, résumé et lien vers le détail.",
        "Un poste fermé doit être retiré ou marqué clairement, sans conserver un bouton actif.",
        "Les offres doivent être partageables via une URL stable et des métadonnées sociales.",
    ]: add_bullet(doc, item)

    add_page_break(doc)
    page_title(doc, "11", "Carrières - détail et candidature")
    add_image(doc, "careers_detail", height=5.35, caption="Maquette 06 - fiche de poste avec panneau de candidature")
    heading(doc, "Structure de la fiche", 2)
    add_text(doc, "Intitulé, contrat, localisation, disponibilité, mission, responsabilités, profil recherché, proposition employeur, documents recommandés et autres opportunités.")
    heading(doc, "Parcours de candidature", 2)
    for item in [
        "Le bouton “Postuler maintenant” ouvre un formulaire accessible ou un canal clairement identifié.",
        "Limiter les pièces obligatoires et indiquer formats, taille maximale et durée de conservation.",
        "Afficher un accusé de réception et préciser le délai indicatif de réponse.",
        "Recueillir le consentement au traitement des données personnelles.",
    ]: add_bullet(doc, item)

    add_page_break(doc)
    page_title(doc, "12", "Contact et conversion")
    add_image(doc, "contact", height=4.15, caption="Maquette 02 - formulaire, coordonnées et carte")
    heading(doc, "Formulaire recommandé", 2)
    add_matrix(doc, ["Champ", "Statut", "Règle"], [
        ("Nom et prénom", "Obligatoire", "Texte, 2 à 80 caractères"),
        ("Société", "Optionnel", "Texte, 2 à 120 caractères"),
        ("E-mail", "Obligatoire", "Validation syntaxique et message explicite"),
        ("Téléphone", "Optionnel", "Format international accepté"),
        ("Service concerné", "Obligatoire", "Liste contrôlée"),
        ("Message", "Obligatoire", "20 à 2 000 caractères"),
        ("Consentement", "Obligatoire", "Case non précochée, lien vers la politique"),
    ], [1.45, 1.05, 3.75])
    callout(doc, "Après envoi", "Afficher une confirmation, envoyer un e-mail de réception, conserver une trace côté équipe et empêcher les doubles soumissions. Ajouter une protection anti-spam respectueuse de l’utilisateur.")

    add_page_break(doc)
    page_title(doc, "13", "Exigences transverses")
    heading(doc, "Responsive", 2)
    for item in [
        "Maquettes de référence à prévoir au minimum pour 1440 px, 768 px et 390 px.",
        "Grilles de trois cartes ramenées à deux puis une colonne, sans perte de contenu.",
        "Images recadrées avec point focal ; titres longs testés sur petits écrans.",
        "Aucune interaction essentielle ne doit dépendre du hover.",
    ]: add_bullet(doc, item)
    heading(doc, "Accessibilité", 2)
    for item in [
        "Contraste conforme WCAG AA, focus visible et navigation complète au clavier.",
        "Un seul H1 par page, ordre logique des titres et liens explicites.",
        "Labels persistants pour les formulaires, erreurs associées aux champs et annonces compatibles lecteur d’écran.",
        "Textes alternatifs pour les images utiles ; attribut vide pour les images décoratives.",
        "Respect de la préférence de réduction des animations.",
    ]: add_bullet(doc, item)
    heading(doc, "Performance et SEO", 2)
    for item in [
        "Images WebP/AVIF responsives, chargement différé sous la ligne de flottaison.",
        "Objectifs : LCP < 2,5 s, INP < 200 ms, CLS < 0,1 sur mobile au 75e percentile.",
        "Titre et description uniques, données structurées Organization et JobPosting si pertinentes.",
        "Plan de site XML, robots.txt, URLs canoniques et partage social configurés.",
    ]: add_bullet(doc, item)

    add_page_break(doc)
    page_title(doc, "14", "Gestion de contenu et données")
    heading(doc, "Contenus administrables", 2)
    add_matrix(doc, ["Type", "Champs principaux", "Cycle de vie"], [
        ("Page", "Titre, slug, sections, SEO, statut", "Brouillon > validation > publication"),
        ("Solution", "Nom, promesse, description, visuel, CTA", "Active / archivée"),
        ("Étude de cas", "Client/secteur, défi, réponse, impact, médias", "Brouillon > validée > publiée"),
        ("Offre d’emploi", "Contrat, lieu, mission, profil, date de clôture", "Ouverte > suspendue > clôturée"),
        ("Partenaire", "Nom, logo, URL, texte alternatif", "Actif / inactif"),
        ("Équipe", "Nom, fonction, bio, photo", "Publié / masqué"),
    ], [1.15, 3.1, 2.0])
    heading(doc, "Gouvernance", 2)
    for item in [
        "Propriétaire éditorial identifié pour chaque type de contenu.",
        "Validation systématique des chiffres, noms, coordonnées et liens avant publication.",
        "Journal des modifications et possibilité de restaurer une version précédente.",
        "Revue trimestrielle des pages institutionnelles et mensuelle des offres d’emploi.",
    ]: add_bullet(doc, item)

    add_page_break(doc)
    page_title(doc, "15", "Recette et critères d’acceptation")
    add_matrix(doc, ["Domaine", "Critère de réussite"], [
        ("Navigation", "Tous les liens mènent à la bonne destination ; état actif et menu mobile validés."),
        ("Contenu", "Aucun texte factice, doublon, faute bloquante ou donnée non vérifiée."),
        ("Formulaires", "Validation, consentement, anti-spam, confirmation et notification testés."),
        ("Responsive", "Aucun débordement horizontal entre 320 px et 1920 px."),
        ("Accessibilité", "Parcours clavier complet, focus visible, labels et contrastes conformes."),
        ("Performance", "Objectifs Core Web Vitals atteints sur pages représentatives."),
        ("SEO", "Métadonnées, canonical, indexation, sitemap et données structurées validés."),
        ("Sécurité", "HTTPS, en-têtes, limitation des soumissions, fichiers contrôlés et secrets protégés."),
        ("Analytics", "Événements CTA, formulaires et candidatures documentés sans données sensibles."),
    ], [1.4, 4.85])
    heading(doc, "Priorités avant développement", 2)
    for item in [
        "P0 - Valider la version canonique de l’accueil et l’arborescence.",
        "P0 - Remplacer les contenus génériques et confirmer les chiffres clés.",
        "P0 - Définir le fonctionnement réel des demandes de démo et candidatures.",
        "P1 - Produire les maquettes tablette/mobile et les états focus/erreur/succès.",
        "P1 - Normaliser le footer, les cartes et les composants de formulaire.",
        "P2 - Préparer le modèle de gestion de contenu et le plan de mesure.",
    ]: add_bullet(doc, item)

    add_page_break(doc)
    page_title(doc, "16", "Inventaire des maquettes sources")
    rows = []
    labels = {
        "home": "Accueil - version principale", "home_alt": "Accueil - alternative",
        "home_classic": "Accueil - classique", "contact": "Contact",
        "solutions": "Solutions - vue principale", "solutions_hover": "Solutions - hover",
        "about": "À propos d’ADC", "former": "Approche - Former", "equiper": "Approche - Équiper",
        "accompagner": "Approche - Accompagner", "careers_list": "Carrières - liste",
        "careers_detail": "Carrières - détail", "cases": "Études de cas", "concept": "Concept hero",
    }
    for key, path in IMAGES.items():
        with Image.open(path) as im:
            dims = f"{im.width} × {im.height} px"
        rows.append((labels[key], path.name, dims))
    add_matrix(doc, ["Écran", "Fichier source", "Dimensions"], rows, [1.75, 3.65, 0.85])
    callout(doc, "Statut", "Cette documentation est fondée sur les quatorze maquettes fournies. Les comportements non visibles ont été formulés comme recommandations fonctionnelles à valider avec l’équipe projet.")
    add_text(doc, "Fin du document", 9, MID, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=18)

    # Metadata
    doc.core_properties.title = "Documentation fonctionnelle et UI - Site web ADC"
    doc.core_properties.subject = "Spécification des maquettes Africa Digital Connect"
    doc.core_properties.author = "Africa Digital Connect"
    doc.core_properties.keywords = "ADC, site web, UX, UI, documentation fonctionnelle"
    doc.core_properties.comments = "Document généré à partir des maquettes fournies."
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
